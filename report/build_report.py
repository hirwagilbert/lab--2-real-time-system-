"""
Builds the lab report (Word) from the code's own results.
    cd <repo root> && python report/build_report.py [--toc toc_pages.json]
Then convert to PDF with LibreOffice (see report/make_report.sh).
"""
from __future__ import annotations

import io
import json
import os
import sys
from contextlib import redirect_stdout

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

from rts_sim import Task, rm_bound, total_utilization
from schedcheck import analyze, load_csv
import partition

BLUE = RGBColor(0x1B, 0x5E, 0x8C)
DARK = RGBColor(0x22, 0x2B, 0x36)
GREY = RGBColor(0x5F, 0x6B, 0x7A)
HEAD_FILL = "1B5E8C"
ZEBRA = "EEF4F8"
REPO = "https://github.com/hirwagilbert/lab--2-real-time-system-"
MEMBERS = [("Janvier NSHIMIYIMANA", "223026679"),
           ("Blandon ASHIMIRWE", "223026675"),
           ("Gilbert HIRWA NTWALI", "223026689")]

TOC_FILE = os.path.join(ROOT, "report", "toc_pages.json")
toc_pages = {}
if "--toc" in sys.argv and os.path.exists(TOC_FILE):
    toc_pages = json.load(open(TOC_FILE))

doc = Document()
HEADINGS = []   # (level, text) in order, for the TOC

# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(11)
st.font.color.rgb = DARK
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

for lvl, size, color, before, after in ((1, 16, BLUE, 18, 8), (2, 13, BLUE, 14, 6), (3, 11.5, DARK, 10, 4)):
    hs = doc.styles[f"Heading {lvl}"]
    hs.font.name = "Calibri"
    hs.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    rf = hs.element.rPr.rFonts
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Calibri")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    hs.font.size = Pt(size)
    hs.font.bold = True
    hs.font.italic = False
    hs.font.color.rgb = color
    hs.paragraph_format.space_before = Pt(before)
    hs.paragraph_format.space_after = Pt(after)
    hs.paragraph_format.keep_with_next = True

cap = doc.styles["Caption"]
cap.font.name = "Calibri"
cap.font.size = Pt(9)
cap.font.italic = True
cap.font.bold = False
cap.font.color.rgb = GREY

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.3)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
CONTENT_W = sec.page_width - sec.left_margin - sec.right_margin


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def borders(table, color="B8C7D3"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def para(text="", bold=False, italic=False, size=None, align=None, color=None, space_after=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_runs(p, text, bold=bold, italic=italic, size=size, color=color)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_runs(p, text, bold=False, italic=False, size=None, color=None):
    """Supports **bold** and _italic_-free simple markup: segments wrapped in ** are bold."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        r = p.add_run(seg)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
    return p


def body(text):
    p = para(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(3)
    if level:
        p.paragraph_format.left_indent = Cm(1.2 + 0.6 * level)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(3)
    return p


def heading(text, level):
    HEADINGS.append((level, text))
    return doc.add_heading(text, level)


fig_no = [0]
tab_no = [0]


def figure(path, caption, width_cm=16.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    fig_no[0] += 1
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figure {fig_no[0]}. "); r.bold = True; r.italic = True
    c.add_run(caption)
    c.paragraph_format.space_after = Pt(10)
    return fig_no[0]


def table(header, rows, widths_cm, caption=None, align=None, font=9.5, bold_first_col=False):
    if caption:
        tab_no[0] += 1
        c = doc.add_paragraph(style="Caption")
        r = c.add_run(f"Table {tab_no[0]}. "); r.bold = True; r.italic = True
        c.add_run(caption)
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_after = Pt(3)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t); cell_margins(t)
    align = align or ["l"] + ["c"] * (len(header) - 1)
    amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
    for i, row in enumerate([header] + rows):
        tr = t.rows[i]
        trPr = tr._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if i == 0:
            th = OxmlElement("w:tblHeader"); trPr.append(th)
        for j, val in enumerate(row):
            cell = tr.cells[j]
            cell.width = Cm(widths_cm[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = amap[align[j]]
            run = p.add_run(str(val))
            run.font.size = Pt(font)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cell, HEAD_FILL)
            else:
                if bold_first_col and j == 0:
                    run.bold = True
                if i % 2 == 0:
                    shade(cell, ZEBRA)
    # grid widths
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[j] * 567)))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def code(text, size=8.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, "D5DEE6"); cell_margins(t, 80, 80, 140, 140)
    c = t.rows[0].cells[0]
    c.width = CONTENT_W
    shade(c, "F4F7FA")
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    lines = text.strip("\n").split("\n")
    for k, ln in enumerate(lines):
        r = p.add_run(ln)
        r.font.name = "DejaVu Sans Mono"
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "DejaVu Sans Mono")
        r.font.size = Pt(size)
        if k < len(lines) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def field(run, instr):
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        run._r.append(e)


def bottom_rule(p, color="1B5E8C", sz=8):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
    bdr.append(b); pPr.append(bdr)


# ---------------------------------------------------------------------------
# Data from the code
# ---------------------------------------------------------------------------
EX1 = [Task("τ1", 1, 4), Task("τ2", 2, 6), Task("τ3", 2, 12)]
EX2 = EX1 + [Task("τ4", 1, 8)]
SETS22 = [("ex_2_2_2", "Section 2.2.2 example"), ("ex_2_3_2", "Section 2.3.2 example"),
          ("harmonic_full", "Harmonic, U = 1"), ("eight_tasks", "Eight tasks"),
          ("constrained", "Constrained deadline"), ("rm_fails_edf_ok", "RM fails, EDF passes"),
          ("overload", "Overload, U > 1")]
REPS22 = {k: analyze(load_csv(f"tasksets/{k}.csv")) for k, _ in SETS22}
with redirect_stdout(io.StringIO()):
    TABLE23, PARTS23 = partition.run_all()


def f3(x):
    return f"{x:.3f}"


# ---------------------------------------------------------------------------
# Cover page
# ---------------------------------------------------------------------------
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
p.paragraph_format.space_before = Pt(6)
p.add_run().add_picture("report/assets/ur_logo.png", width=Cm(5.2))

para("UNIVERSITY OF RWANDA", bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, space_after=2)
para("Gako Campus", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, space_after=2)
para("Department of Computer Engineering", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, space_after=2)
para("Year 4  ·  Semester I", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, space_after=14)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14); bottom_rule(p)

para("Module: Real-Time Systems", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, space_after=6)
para("Unit 2 Lab Report", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=4)
para("Real-Time Scheduling Mini Projects 2.1 – 2.3", bold=True, size=19,
     align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, space_after=4)
para("RM, EDF and Priority Inheritance Simulation · RM versus EDF Schedulability Tool · "
     "Partitioned Multiprocessor Scheduling", italic=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER,
     color=GREY, space_after=18)

p = para("GROUP ASSIGNMENT", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, space_after=8)

t = doc.add_table(rows=1 + len(MEMBERS), cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(t); cell_margins(t, 70, 70, 120, 120)
for i, row in enumerate([("No.", "Full name", "Registration number")] +
                        [(str(k + 1), n, r) for k, (n, r) in enumerate(MEMBERS)]):
    for j, val in enumerate(row):
        c = t.rows[i].cells[j]
        c.width = Cm([1.4, 7.0, 5.0][j])
        pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = pp.add_run(val); r.font.size = Pt(11)
        if i == 0:
            r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, HEAD_FILL)
        elif i % 2 == 0:
            shade(c, ZEBRA)
for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
    gc.set(qn("w:w"), str(int([1.4, 7.0, 5.0][j] * 567)))

para(space_after=18)
para("Source code repository", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=0)
para(REPO, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, space_after=10)
para("Submitted: October 2026", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, space_after=0)

# New section for body (header/footer only on body pages)
body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
body_sec.different_first_page_header_footer = False
doc.sections[0].different_first_page_header_footer = True
for s in (body_sec,):
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    hp = s.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("Real-Time Systems · Unit 2 Mini Projects"); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    hp.add_run("\t\t")
    r = hp.add_run("University of Rwanda – Gako Campus"); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    bottom_rule(hp, "B8C7D3", 4)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("Page "); r.font.size = Pt(9); r.font.color.rgb = GREY
    r = fp.add_run(); r.font.size = Pt(9); r.font.color.rgb = GREY
    field(r, "PAGE")
    gp = s.footer.add_paragraph()
    gp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gp.paragraph_format.space_after = Pt(0)
    r = gp.add_run("GitHub: " + REPO); r.font.size = Pt(8); r.font.color.rgb = BLUE
# First section (cover) has empty header/footer
doc.sections[0].first_page_header.paragraphs[0].text = ""
doc.sections[0].first_page_footer.paragraphs[0].text = ""

# ---------------------------------------------------------------------------
# Table of contents (static, page numbers filled in on the second pass)
# ---------------------------------------------------------------------------
TOC_PLAN = []  # filled after headings are known: we pre-declare the heading list below


def toc_block(entries):
    p = para("Table of Contents", bold=True, size=16, color=BLUE, space_after=10)
    for lvl, text in entries:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_after = Pt(3 if lvl == 1 else 1)
        tp.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.7)
        tp.paragraph_format.tab_stops.add_tab_stop(CONTENT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = tp.add_run(text); r.bold = (lvl == 1); r.font.size = Pt(11 if lvl == 1 else 10.5)
        r = tp.add_run("\t" + str(toc_pages.get(text, "")))
        r.bold = (lvl == 1); r.font.size = Pt(11 if lvl == 1 else 10.5)


# The TOC lists level-1 and level-2 headings; these are declared here so the TOC
# can precede the content. They must match the headings written below.
TOC_ENTRIES = [
    (1, "1. Introduction"),
    (2, "1.1 Objectives"), (2, "1.2 Environment and repository layout"),
    (1, "2. Mini Project 2.1 – Simulating RM, EDF and Priority Inheritance"),
    (2, "2.1 Task model and utilization tests"), (2, "2.2 RM and EDF simulation"),
    (2, "2.3 Response Time Analysis of τ4"), (2, "2.4 Priority inversion and the PIP fix"),
    (2, "2.5 Why the simulated PIP blocking matches the bound"), (2, "2.6 Lab tasks"),
    (1, "3. Mini Project 2.2 – RM versus EDF Schedulability Comparison Tool"),
    (2, "3.1 Tool design"), (2, "3.2 Results on seven task sets"), (2, "3.3 Written report"),
    (2, "3.4 Reflection: when the bound fails but RTA succeeds"),
    (1, "4. Mini Project 2.3 – Partitioned Multiprocessor Scheduling Simulator"),
    (2, "4.1 Design"), (2, "4.2 Test sets and comparison table"), (2, "4.3 Per-core schedule"),
    (2, "4.4 Written analysis"),
    (1, "5. Conclusion"), (1, "References"), (1, "Appendix A – Reproducing the results"),
]
toc_block(TOC_ENTRIES)
page_break()

# ---------------------------------------------------------------------------
# 1. Introduction
# ---------------------------------------------------------------------------
heading("1. Introduction", 1)
body("This report documents our group's solution to the three mini projects of the Unit 2 lab manual of the "
     "Real-Time Systems module. All three projects were implemented as pure-Python discrete-event (tick-based) "
     "simulators – no RTOS, board or C compiler is needed – and every number, table and chart in this report "
     "was produced by the code in our repository, so the results can be reproduced exactly by re-running the "
     "scripts listed in Appendix A.")
heading("1.1 Objectives", 2)
bullet("**Mini Project 2.1** – build a simulator for Rate Monotonic (RM) and Earliest Deadline First (EDF), apply "
       "the Liu & Layland and EDF U ≤ 1 tests, run Response Time Analysis (RTA), and reproduce priority inversion "
       "together with its Priority Inheritance Protocol (PIP) fix.")
bullet("**Mini Project 2.2** – generalise the simulator into a tool that takes any periodic task set (from a list or "
       "a CSV file) and reports RM/DM and EDF schedulability with an RTA fallback and a two-panel comparison chart.")
bullet("**Mini Project 2.3** – partition task sets over m cores with First-Fit, First-Fit Decreasing, Best-Fit "
       "Decreasing and Worst-Fit, using the per-core Liu & Layland bound as the admission test.")
heading("1.2 Environment and repository layout", 2)
body("Python 3.10+ and matplotlib (charts drawn with broken_barh) are the only requirements; pandas was not needed. "
     "Time is discrete: one tick equals one time unit (1 ms in the unit's examples), all jobs of a task are released "
     "synchronously at multiples of Tᵢ, and a job that misses its deadline is recorded and allowed to finish late. "
     f"The complete code is available at {REPO}.")
table(["File / folder", "Purpose"],
      [["rts_sim.py", "Shared core: Task model, utilization tests, RTA, RM/DM/EDF/LLF tick simulators, "
                      "mutex simulator with/without PIP, Gantt helpers"],
       ["mp21_run.py", "Runs every procedure step and lab task of Mini Project 2.1"],
       ["schedcheck.py", "Mini Project 2.2 command-line tool (CSV or Python list input)"],
       ["tasksets/*.csv", "The seven task sets analysed in Mini Project 2.2"],
       ["partition.py", "Mini Project 2.3: FF, FFD, BFD, WF partitioning and per-core simulation"],
       ["tests/test_all.py", "Hand-computed checks (bounds, RTA, PIP blocking, verdicts, partitions)"],
       ["figures/, results/", "Generated PNG charts and the printed text output of each script"],
       ["report/", "This report (PDF) and the script that builds it"]],
      [4.2, 12.2], "Repository contents.", align=["l", "l"])

# ---------------------------------------------------------------------------
# 2. Mini Project 2.1
# ---------------------------------------------------------------------------
page_break()
heading("2. Mini Project 2.1 – Simulating RM, EDF and Priority Inheritance", 1)
heading("2.1 Task model and utilization tests", 2)
body("Each periodic task τᵢ is a Task dataclass holding name, Cᵢ, Tᵢ and Dᵢ (Dᵢ defaults to Tᵢ, the implicit-deadline "
     "model of Section 2.1.2) with the derived utilization uᵢ = Cᵢ/Tᵢ. The functions total_utilization(tasks) and "
     "rm_bound(n) = n(2^(1/n) − 1) implement Section 2.2.1, and the EDF test checks U ≤ 1 (Section 2.3.1). "
     "Note that the lab manual's printed formula \"2 * (1/n)\" is a typesetting slip for the power 2^(1/n); our "
     "implementation uses the power, which is what gives the expected value 0.780 for n = 3.")
rows = []
for label, ts in (("Section 2.2.2 (3 tasks)", EX1), ("Section 2.3.2 (+ τ4: C=1, T=8)", EX2)):
    U, b = total_utilization(ts), rm_bound(len(ts))
    rows.append([label, " ".join(f"{t.name}({t.C},{t.T})" for t in ts), f"{U:.3f}", f"{b:.3f}",
                 "PASS" if U <= b else "FAIL (inconclusive)", "PASS" if U <= 1 else "FAIL"])
table(["Example", "Tasks (C, T)", "U", "L&L bound", "RM bound test", "EDF U ≤ 1"], rows,
      [3.6, 4.6, 1.4, 2.0, 2.8, 2.0], "Utilization tests reproduced for the unit's two worked examples.",
      font=9)
body("Both results match the unit exactly: the three-task set (U = 0.750 ≤ 0.780) is guaranteed RM-schedulable by the "
     "sufficient test, while adding τ4 raises U to 0.875, above rm_bound(4) = 0.757. The RM test is then "
     "**inconclusive** (it is only sufficient), whereas EDF's exact test U ≤ 1 still guarantees schedulability.")

heading("2.2 RM and EDF simulation", 2)
body("simulate_rm(tasks, sim_time) steps through time one tick at a time. At every tick t it releases a new job of "
     "τᵢ whenever t is a multiple of Tᵢ, records any job whose absolute deadline has passed unfinished, and runs the "
     "ready job of the highest-priority task (shortest period), preempting the running job if necessary. "
     "simulate_edf is identical except that the selection key is the job's absolute deadline; ties on deadline "
     "are broken in favour of the task listed first (the shorter-period task). Both return the per-tick timeline, "
     "deadline misses, response times and context-switch counts. Figure 1 shows the four-task set over one "
     "hyperperiod (LCM(4, 6, 12, 8) = 24 ms).")
figure("figures/mp21_rm_vs_edf.png", "RM (top) and EDF (bottom) schedules of the four-task set over one hyperperiod. "
       "Triangles mark job releases; grey hatching is idle time.")
rmr = REPS22["ex_2_3_2"]
table(["Policy", "Deadline misses", "Context switches", "Worst R(τ1)", "Worst R(τ2)", "Worst R(τ3)", "Worst R(τ4)", "Idle"],
      [["RM", "0", "15", "1", "3", "11", "4", "3"], ["EDF", "0", "15", "1", "3", "10", "6", "3"]],
      [1.6, 2.2, 2.3, 1.9, 1.9, 1.9, 1.9, 1.2], "Simulated metrics over one hyperperiod (24 ms).", font=9)
body("**Observations.** τ1 executes in exactly the same slots in both charts, because it always has both the shortest "
     "period and the earliest deadline. The schedules diverge at t = 9, 10, 17 and 21, where the jobs of τ3 and τ4 "
     "interleave: RM always favours τ4 (period 8) over τ3 (period 12), whereas EDF runs whichever job has the nearer "
     "absolute deadline – at t = 9 the first job of τ3 (deadline 12) is served before the second job of τ4 (deadline 16). "
     "EDF therefore lowers τ3's worst-case response time from 11 to 10 ms at the cost of τ4 (4 → 6 ms), i.e. it "
     "spreads the slack according to urgency instead of according to a static rate. Total busy time is identical "
     "(21 of 24 ms, U = 0.875) because both algorithms are work-conserving, and both meet every deadline – "
     "something the RM utilization bound alone could not promise.")

heading("2.3 Response Time Analysis of τ4", 2)
body("Because the utilization bound was inconclusive, we applied the exact RTA recurrence of Section 2.5.2, "
     "Rᵢ = Cᵢ + Σⱼ∈hp(i) ⌈Rᵢ/Tⱼ⌉·Cⱼ, iterating from Rᵢ = Cᵢ until convergence or until Rᵢ exceeds Dᵢ. Under RM the "
     "priority order is τ1 (T=4) > τ2 (T=6) > τ4 (T=8) > τ3 (T=12); τ4 is therefore third, and τ3 – not τ4 – is "
     "the lowest-priority task, so we analysed both.")
code("""τ4: hp = {τ1, τ2}       R⁰ = 1
     R¹ = 1 + ⌈1/4⌉·1 + ⌈1/6⌉·2 = 4
     R² = 1 + ⌈4/4⌉·1 + ⌈4/6⌉·2 = 4   → converged, R4 = 4 ≤ D4 = 8   ✔
τ3: hp = {τ1, τ2, τ4}   R: 2 → 6 → 7 → 9 → 11 → 11  → R3 = 11 ≤ D3 = 12  ✔""")
table(["Task", "RM priority", "C", "T = D", "R (RTA)", "Iterations", "Verdict"],
      [["τ1", "1", "1", "4", "1", "1", "OK"], ["τ2", "2", "2", "6", "3", "2", "OK"],
       ["τ4", "3", "1", "8", "4", "2", "OK"], ["τ3", "4", "2", "12", "11", "5", "OK"]],
      [1.6, 2.3, 1.3, 1.8, 2.0, 2.0, 2.0], "RTA results under RM for the four-task set.")
body("RTA gives a definitive answer where the utilization bound could not: R4 = 4 ≤ 8 and every Rᵢ ≤ Dᵢ, so the "
     "set **is RM-schedulable**. The analytical values agree exactly with the worst response times observed in "
     "the simulation (Table 3), which validates both the simulator and the RTA implementation.")

heading("2.4 Priority inversion and the PIP fix", 2)
body("For Section 2.6.1 we built a second simulator, simulate_resources, in which each job is a list of "
     "segments – normal execution or a critical section on a named mutex. The scenario is: low-priority L "
     "takes mutex S at t = 0 and needs 6 ticks inside it; high-priority H arrives at t = 1 and needs S for 2 ticks; "
     "medium-priority M arrives at t = 2 and computes for 4 ticks without touching S. With protocol = 'none' the "
     "scheduler simply runs the highest-priority job that is not blocked. With protocol = 'pip', every tick the "
     "holder of a mutex inherits the highest priority among the jobs blocked on it (applied transitively); because "
     "the base priority is never overwritten, it is automatically restored as soon as the mutex is released.")
figure("figures/mp21_inversion_vs_pip.png", "Priority inversion without a protocol (top) and with the Priority "
       "Inheritance Protocol (bottom). Hatched = inside the critical section; cross-hatched = L running at H's "
       "inherited priority; red band = H ready but blocked.")
body("Without a protocol, M preempts L at t = 2 even though H is waiting for L, so H is delayed by the rest of L's "
     "critical section **and** by all of M's execution; it completes at t = 12. With PIP, L runs at H's priority from "
     "t = 1, M cannot preempt it, L releases S at t = 6, H completes at t = 8 and M runs afterwards. This is the "
     "simulated analogue of the Mars Pathfinder fix.")
table(["Scenario", "Protocol", "H finishes", "Response R_H", "H blocked (ticks)", "PIP bound"],
      [["L, M, H", "None", "12", "11", "9 = 5 (CS of L) + 4 (M)", "—"],
       ["L, M, H", "PIP", "8", "7", "5", "≤ 6"],
       ["L, M, M2, H", "None", "15", "14", "12 = 5 + 4 + 3", "—"],
       ["L, M, M2, H", "PIP", "8", "7", "5", "≤ 6"]],
      [2.6, 1.7, 2.0, 2.3, 4.6, 1.8], "Measured blocking of H (arrival t = 1, C_H = 2).")

heading("2.5 Why the simulated PIP blocking matches the bound", 2)
body("Section 2.7.1 states that under the Priority Inheritance Protocol a high-priority job can be blocked for at "
     "most one critical section of each lower-priority job that shares a resource with it. In our scenario only one "
     "lower-priority job, L, shares the mutex S with H, and L's critical section lasts 6 ticks, so the bound on "
     "H's blocking is 6 ticks. The simulator measured 5 ticks, which respects the bound, and the difference is "
     "fully explained by the timeline: L had already executed one tick of its critical section before H arrived "
     "at t = 1, so only the remaining 5 ticks could block H. The decisive point is what happens at t = 2. Without "
     "a protocol, M has a higher priority than L, preempts it, and H must additionally wait for all 4 ticks of M, "
     "an unrelated task; the blocking grows to 9 ticks and, when we added a second medium task M2, to 12 ticks. "
     "Inversion of this kind is unbounded because it depends on how much medium-priority work happens to arrive. "
     "Under PIP, L inherits H's priority the moment H blocks, so neither M nor M2 can preempt it, and as soon as "
     "L releases S its base priority is restored and H runs immediately. H's blocking is therefore limited to the "
     "remainder of one critical section of L – 5 ticks with one or two medium tasks alike – exactly as the "
     "one-critical-section-per-lower-priority-task bound predicts.")

heading("2.6 Lab tasks", 2)
heading("Lab task 1 – Deadline Monotonic", 3)
body("simulate_dm orders priorities by relative deadline. On the implicit-deadline four-task set its timeline is "
     "identical, tick for tick, to RM (Dᵢ = Tᵢ makes the two orders the same). We then shortened τ3's deadline to "
     "D3 = 3 (τ1 (1,4,4), τ2 (2,6,6), τ3 (2,12,3)). RM still gives τ3 the lowest priority and it misses its first "
     "deadline (RTA: R3 = 5 > 3), whereas DM promotes τ3 to the top and all tasks meet their deadlines "
     "(R3 = 2, R1 = 3, R2 = 6). The schedule changes only when some Dᵢ ≠ Tᵢ.")
figure("figures/mp21_lab_dm.png", "RM versus Deadline Monotonic on the constrained-deadline set; the red X is τ3's "
       "deadline miss under RM.", 15.5)
heading("Lab task 2 – Bound sweep", 3)
body("Sweeping n = 1…10 gives 1.000, 0.828, 0.780, 0.757, 0.743, 0.735, 0.729, 0.724, 0.721 and 0.718; the curve "
     "decreases monotonically towards ln 2 ≈ 0.693, while the EDF bound stays at 1.0 (Figure 4). The gap between the "
     "two curves is the utilization that RM's sufficient test may waste.")
figure("figures/mp21_lab_bound_sweep.png", "Liu & Layland RM bound versus the constant EDF bound.", 12.5)
heading("Lab task 3 – Least Laxity First", 3)
body("simulate_llf selects the job with the smallest laxity d − t − remaining. On the light four-task set all jobs "
     "are 1–2 ticks long, so LLF has no room to thrash and produced the same 15 context switches as RM and EDF. "
     "On a heavier set with longer jobs, a (4, 10), b (5, 12), c (2, 15), U = 0.95, the difference is clear over "
     "one hyperperiod (60 ms): RM misses 3 deadlines with 19 context switches, EDF meets every deadline with 15, "
     "and LLF also meets every deadline but needs **31 context switches (17 preemptions)**. Because the running "
     "job's laxity stays constant while the waiting jobs' laxity drops every tick, two jobs with similar laxity "
     "keep overtaking each other – the thrashing visible in the a/b alternation of Figure 5.")
figure("figures/mp21_lab_llf_heavy.png", "RM, EDF and LLF on the heavier set (U = 0.95) over one hyperperiod.")
heading("Lab task 4 – A second medium-priority task", 3)
body("Adding M2 (priority equal to M, arrival t = 3, 3 ticks) during M's execution increases H's unprotected blocking "
     "from 9 to 12 ticks, while the PIP-protected blocking stays at 5 ticks (Table 5 and Figure 6).")
figure("figures/mp21_inversion_vs_pip_2medium.png", "Inversion with two medium-priority tasks: unprotected "
       "blocking grows, PIP blocking does not.")

# ---------------------------------------------------------------------------
# 3. Mini Project 2.2
# ---------------------------------------------------------------------------
page_break()
heading("3. Mini Project 2.2 – RM versus EDF Schedulability Comparison Tool", 1)
heading("3.1 Tool design", 2)
body("schedcheck.py reuses the simulators of rts_sim.py and adds an analysis layer. The tool accepts either a Python "
     "list of Task objects (analyze(tasks)) or one or more CSV files with columns name, C, T and an optional D "
     "(blank D defaults to T). Its decision layer follows Section 2.5.2:")
numbered("Compute uᵢ, U and the hyperperiod (LCM of all Tᵢ).")
numbered("If every Dᵢ = Tᵢ, use RM priorities and test U ≤ n(2^(1/n) − 1). If any Dᵢ < Tᵢ, switch to Deadline "
         "Monotonic priorities and apply the same bound to the density Σ Cᵢ/Dᵢ, which is a valid sufficient test for DM.")
numbered("If the bound passes, report schedulable directly; if it fails, run RTA for every task in priority order "
         "and report Rᵢ, the iteration count and a per-task verdict.")
numbered("EDF verdict: U ≤ 1 (exact for implicit deadlines). For constrained deadlines U ≤ 1 is only necessary, so the "
         "tool additionally runs the exact processor-demand test (demand bound function ≤ L at every absolute deadline "
         "up to H + max D).")
numbered("Cross-check both verdicts by simulating one hyperperiod (capped at 500 time units) and draw a two-panel "
         "Gantt chart, RM/DM on top and EDF below, with distinct colours per task, idle time and red crosses at misses.")
body("Robustness: a set with U > 1 is reported as infeasible under any algorithm and still charted (the late jobs and "
     "misses are visible); RTA stops as soon as Rᵢ > Dᵢ, so it cannot loop forever; malformed CSV rows, non-integer "
     "values, duplicate names and D > T produce a clear error message and a non-zero exit code instead of a crash.")
code("""$ python schedcheck.py tasksets/ex_2_3_2.csv
n = 4   U = 0.8750   hyperperiod = 24
Liu & Layland bound n(2^(1/n)-1) = 0.7568   U <= bound ? FAIL
Response Time Analysis (RM priority order):
   Task     prio   R_i   D_i  iter  verdict
   t1          1     1     4     1  OK
   t2          2     3     6     2  OK
   t4          3     4     8     2  OK
   t3          4    11    12     5  OK
RM verdict : SCHEDULABLE (exact RTA)
EDF verdict: SCHEDULABLE (U <= 1)
Simulation over 24 units: RM misses = 0, EDF misses = 0""", 8)

heading("3.2 Results on seven task sets", 2)
rows = []
for k, label in SETS22:
    r = REPS22[k]
    tset = ", ".join(f"({t['C']},{t['T']}" + (f",{t['D']})" if t['D'] != t['T'] else ")") for t in r["tasks"])
    rows.append([label, str(r["n"]), f"{r['U']:.3f}", f"{r['bound']:.3f}",
                 ("PASS" if r["bound_pass"] else "FAIL"),
                 r["fp_policy"] + ": " + ("OK" if r["fp_ok"] else "MISS") + ("" if r["bound_pass"] else " (RTA)"),
                 "OK" if r["edf_ok"] else ("INFEASIBLE" if r["infeasible"] else "MISS"),
                 f"{len(r['sim_fp'].misses)} / {len(r['sim_edf'].misses)}"])
table(["Task set", "n", "U", "L&L bound", "Bound test", "Fixed priority", "EDF", "Sim. misses FP / EDF"], rows,
      [3.6, 0.8, 1.3, 1.6, 1.7, 2.7, 2.1, 2.6], "Summary of the tool's verdicts. Bound test uses density ΣC/D for "
      "the constrained set.", font=8.5)
r8 = REPS22["eight_tasks"]
table(["Set", "Tasks (C, T[, D])"],
      [[label, ", ".join(f"{t['name']}({t['C']},{t['T']}" + (f",{t['D']})" if t['D'] != t['T'] else ")")
                         for t in REPS22[k]["tasks"])] for k, label in SETS22],
      [3.6, 12.8], "Task-set definitions (files in tasksets/).", align=["l", "l"], font=8.5)

heading("3.3 Written report", 2)
body("We validated the tool on the two worked examples of the unit and on five task sets of our own, chosen so that "
     "every branch of the decision layer is exercised. In every case the analytical verdict agreed with the simulated "
     "schedule over one hyperperiod (last column of Table 6), which gives us confidence that the utilization tests, the "
     "RTA fallback and the simulators are mutually consistent.")
body("**Worked examples.** The Section 2.2.2 set (U = 0.750) passes the Liu & Layland bound of 0.780, so the tool "
     "reports RM-schedulable immediately without running RTA. The Section 2.3.2 set (U = 0.875) fails the bound of "
     "0.757 and the tool falls back to RTA, which returns R = 1, 3, 4 and 11 for τ1, τ2, τ4, τ3 – identical to our "
     "hand computation in Section 2.3 – so the set is RM-schedulable. EDF passes in both cases.")
figure("figures/mp22_ex_2_3_2.png", "Tool output chart for the Section 2.3.2 set: bound fails, RTA proves RM "
       "schedulability; both schedules meet every deadline.")
body("**Eight tasks.** With periods 10…60 ms and U = 0.842, the bound for n = 8 has already fallen to 0.724, close to "
     "its asymptotic ln 2 limit. The sufficient test fails by a wide margin, yet RTA shows that even the lowest-priority "
     "task (C = 6, T = 60) has R = 40 ≤ 60. This illustrates how pessimistic the bound becomes for larger task sets: "
     "more than 11 % of the processor would have been declared unusable.")
figure("figures/mp22_eight_tasks.png", "Eight-task set (U = 0.842) over its 120 ms hyperperiod; RM and EDF both "
       "meet all deadlines.")
body("**Constrained deadline.** For τ1 (1,4,4), τ2 (2,6,6), τ3 (2,12,3) the tool detects Dᵢ < Tᵢ and switches to "
     "Deadline Monotonic ordering. The density 1.25 exceeds the bound, but RTA under DM finds R = 2, 3, 6 against "
     "deadlines 3, 4, 6, so the set is schedulable. This set also shows why U ≤ 1 is not an exact EDF test once "
     "deadlines are constrained: the density is above 1 while U is only 0.75, and the exact processor-demand test "
     "is what confirms that EDF meets every deadline.")
body("**RM fails, EDF succeeds.** For a (4,10), b (5,12), c (2,15) with U = 0.95, RTA proves that task c misses its "
     "deadline under RM (its response time exceeds 15 ms) and the simulation shows three misses in 60 ms, while EDF "
     "schedules the set with no misses. This is the classic case in which the dynamic-priority algorithm uses the "
     "processor more fully than any static assignment.")
figure("figures/mp22_rm_fails_edf_ok.png", "U = 0.95: three deadline misses for task c under RM (red crosses), "
       "none under EDF.")
body("**Overload.** The set (2,4), (3,6), (2,8) has U = 1.25. The tool reports it as infeasible on one processor under "
     "any algorithm, RTA marks τ2 and τ3 as missing, and the chart still renders, showing deadline misses under both "
     "policies instead of crashing. EDF's misses cascade once the processor is overloaded, which is the well-known "
     "domino effect that makes EDF less predictable than RM under overload.")
figure("figures/mp22_overload.png", "Overloaded set (U = 1.25): the tool reports infeasibility and shows the misses.")
body("**Harmonic set at full utilization.** Finally, (1,4), (2,8), (4,16), (4,16) has U = 1.0, far above the bound of "
     "0.757, yet RTA gives R = 1, 3, 8, 16 and the simulation is miss-free. When periods are harmonic RM can reach "
     "100 % utilization, so the Liu & Layland bound should never be read as a necessary condition.")

heading("3.4 Reflection: when the bound fails but RTA succeeds", 2)
body("The four-task set of Section 2.3.2 is the clearest example of why exact tests matter. Its utilization of 0.875 is "
     "above the Liu & Layland bound of 0.757, so a designer who relied only on the bound would conclude that RM cannot "
     "be trusted and might add hardware, lengthen periods or switch to EDF unnecessarily. Response Time Analysis tells "
     "a different story: the lowest-priority task τ3 completes after at most 11 ms against a 12 ms deadline, τ4 after "
     "4 ms against 8 ms, and the simulation over the 24 ms hyperperiod confirms these values to the tick. The bound is "
     "derived for the worst possible combination of periods, so it is safe but often pessimistic; the more tasks there "
     "are, the more capacity it gives away, approaching 31 % as n grows. The harmonic set we built, schedulable by RM at "
     "a full 100 % utilization, makes the same point even more strongly. In practice the bound is a fast first filter, "
     "but only an exact test such as RTA can certify a set that fails it, and certifying it lets the same processor carry "
     "considerably more work.")

# ---------------------------------------------------------------------------
# 4. Mini Project 2.3
# ---------------------------------------------------------------------------
page_break()
heading("4. Mini Project 2.3 – Partitioned Multiprocessor Scheduling Simulator", 1)
heading("4.1 Design", 2)
body("partition.py turns the m-processor problem into m independent uniprocessor problems plus a bin-packing "
     "problem (Section 2.8.1). The \"does it fit\" predicate is the Liu & Layland bound recomputed for the number of "
     "tasks the core would hold after the placement – never a flat 0.69 or 1.0 threshold:")
code("core k admits τi  ⇔  U_k + u_i ≤ rm_bound(n_k + 1)    (n_k = tasks on core k)")
body("All four heuristics share the signature heuristic(tasks, m) → Partition, which records the task-to-core "
     "mapping and any task that could not be placed. Remaining capacity is defined as the slack left under the bound "
     "after the placement, rm_bound(n_k + 1) − (U_k + u_i).")
table(["Heuristic", "Task order", "Core chosen among admissible cores"],
      [["First-Fit (FF)", "input order", "the first (lowest-index) core"],
       ["First-Fit Decreasing (FFD)", "decreasing uᵢ", "the first core"],
       ["Best-Fit Decreasing (BFD)", "decreasing uᵢ", "the core with the least remaining capacity"],
       ["Worst-Fit (WF)", "input order", "the core with the most remaining capacity"]],
      [5.0, 3.2, 8.2], "The four bin-packing heuristics.", align=["l", "l", "l"])

heading("4.2 Test sets and comparison table", 2)
body("We used m = 4 and five task sets of increasing difficulty. Periods are drawn from {5, 10, 20, 40} ms so that the "
     "per-core hyperperiods stay short enough to simulate and plot.")
defs = []
for s, ts in partition.TEST_SETS.items():
    U = total_utilization(ts)
    us = sorted({round(t.U, 3) for t in ts})
    defs.append([s.replace("_", " "), str(len(ts)), f"{U:.3f}", f"{100 * U / 4:.1f} %",
                 ", ".join(f"{u:g}" for u in us)])
table(["Set", "n", "Total U", "U / m", "Distinct uᵢ values"], defs, [4.4, 1.0, 1.8, 1.8, 7.4],
      "Partitioning test sets (m = 4).", align=["l", "c", "c", "c", "l"], font=9)
rows = []
for r in TABLE23:
    rows.append([r["set"].replace("_", " "), r["heuristic"], "OK" if r["success"] else "FAIL",
                 f"{r['cores_used']}/4", " / ".join(f"{x:.2f}" for x in r["loads"]),
                 f"{r['placed_U']:.3f}", f"{r['pct_of_m']:.1f} %", str(r["unplaced"]),
                 str(r["cores_needed"]), str(r["lower_bound"])])
table(["Set", "Heur.", "Result", "Cores", "Per-core U (core 0–3)", "Placed U", "% of m", "Unpl.",
       "Need*", "⌈U⌉"], rows, [3.0, 1.1, 1.3, 1.2, 3.6, 1.5, 1.5, 1.0, 1.1, 1.0],
      "Comparison of the heuristics. *Need = cores the heuristic uses when given unlimited cores.", font=8)
figure("figures/mp23_summary.png", "Placed utilization per heuristic, stacked by core; the dashed line is the "
       "task set's total utilization, so a bar below it means some tasks were not placed.")

heading("4.3 Per-core schedule", 2)
body("For the FFD partition of set A we ran the uniprocessor RM simulator independently on each core over one "
     "hyperperiod (40 ms). Every core meets every deadline (Figure 12); the largest worst-case response time, 28 ms "
     "for a10 (T = 40 ms) on core 1, is still well inside its deadline. Each core is loaded to at most 0.75, below its "
     "own bound (0.757 for four tasks), as the admission test guarantees.")
figure("figures/mp23_per_core_gantt_FFD_A.png", "Per-core RM schedules of the FFD partition of set A; "
       "each row is one core, grey is idle time; no deadline is missed.")

heading("4.4 Written analysis", 2)
body("Section 2.8.1 reduces partitioned scheduling to m uniprocessor problems plus an NP-hard bin-packing problem, and "
     "Section 2.8.3 notes that practical systems therefore use heuristics with a per-core schedulability test as the "
     "fit predicate. Our results show where this works and where it breaks.")
body("**Why FFD outperforms FF.** Sorting by decreasing utilization places the hardest items first, while every core "
     "is still empty and has its full capacity. Set D makes the effect visible: with four small tasks (u = 0.2) listed "
     "before four large ones (u = 0.6), FF stacks three small tasks on core 0 (load 0.6), so no large task can join it "
     "any more (0.6 + 0.6 > rm_bound(4)). The last large task then finds no admissible core and FF fails, placing "
     "only 65 % of m. FFD places one large task per core first and then adds one small task to each, reaching exactly "
     "0.8 per core (just under the two-task bound of 0.828) and succeeding with the optimal four cores. On sets A and "
     "B, where the input order happened to be favourable, FF and FFD both succeed, but FFD packs cores 0 and 1 to "
     "0.75 and leaves core 3 almost empty, i.e. it consolidates load and keeps a whole core free for future tasks. BFD "
     "behaved like FFD on all our sets.")
body("**Why WF is weaker.** Worst-Fit spreads load evenly, which looks attractive (set A ends with balanced loads "
     "0.70, 0.68, 0.50, 0.55), but it destroys the large free gaps that big tasks need. In set B the small tasks are "
     "spread over all four cores before the u = 0.65 task arrives; by then every core already carries enough load "
     "that adding 0.65 would violate its bound, so WF fails while FF, FFD and BFD succeed. With an unlimited supply of "
     "cores WF always opens a new, empty core, so it would use one core per task (14 cores for set A).")
body("**The adversarial set.** Set E contains five tasks of u = 0.55 each. Its total utilization is 2.75, only 69 % of "
     "the four-core capacity, yet every heuristic fails. The reason is purely the bin-packing limitation of Section "
     "2.8.1: any two of these tasks together need 1.10, which exceeds the two-task bound of 0.828 and even exceeds a "
     "single processor's absolute capacity of 1.0. Each core can therefore hold only one task, five tasks need five "
     "cores, and no ordering or choice rule can change this. By hand: ⌈2.75⌉ = 3 cores would suffice if work were divisible, "
     "but partitioning cannot split tasks. Set C makes the same point with smaller items: ten tasks of u = 0.375 "
     "(U = 3.75, 94 % of m) fit at most two per core (three would need 1.125), so only eight are placed and 25 % of the "
     "platform is unusable.")
body("**Distance from the optimum.** The trivial lower bound ⌈U⌉ is optimistic because it assumes a core can be "
     "filled to 1.0. Under RM admission no core holding three or more tasks can exceed 0.78, and in set A a core with "
     "at most two tasks holds at most 0.4, so three cores can carry at most about 2.34 < 2.425. Four cores are "
     "therefore genuinely needed, and FFD/BFD's result is optimal under this admission test; in set D FFD also meets "
     "the lower bound of four cores exactly. Capacity is lost to the per-core bound itself and to fragmentation; the "
     "adversarial sets show that fragmentation can dominate even at low total utilization.")

# ---------------------------------------------------------------------------
# 5. Conclusion, references, appendix
# ---------------------------------------------------------------------------
heading("5. Conclusion", 1)
body("The three mini projects reproduced every quantitative result of Unit 2 and extended them. Our simulator confirmed "
     "the utilization results of Sections 2.2.2 and 2.3.2 exactly, showed RM and EDF diverging where τ3 and τ4 "
     "interleave, and confirmed with RTA that the four-task set is RM-schedulable despite failing the bound. The "
     "priority-inversion experiment showed blocking that grows with every medium-priority task (9 → 12 ticks) and "
     "PIP holding it at 5 ticks, within the one-critical-section bound. The comparison tool turned these pieces into a "
     "reusable analyser that handles implicit and constrained deadlines and overloaded sets gracefully. Finally, the "
     "partitioning simulator showed that FFD and BFD outperform FF and WF, and that a set using only 69 % of a "
     "four-core platform can still be impossible to partition – the central weakness of partitioned scheduling.")

heading("References", 1)
refs = [
    "C. L. Liu and J. W. Layland, \"Scheduling algorithms for multiprogramming in a hard-real-time environment,\" "
    "Journal of the ACM, vol. 20, no. 1, pp. 46–61, 1973.",
    "M. Joseph and P. Pandya, \"Finding response times in a real-time system,\" The Computer Journal, vol. 29, no. 5, "
    "pp. 390–395, 1986.",
    "L. Sha, R. Rajkumar and J. P. Lehoczky, \"Priority inheritance protocols: an approach to real-time "
    "synchronization,\" IEEE Transactions on Computers, vol. 39, no. 9, pp. 1175–1185, 1990.",
    "J. Y.-T. Leung and J. Whitehead, \"On the complexity of fixed-priority scheduling of periodic, real-time tasks,\" "
    "Performance Evaluation, vol. 2, no. 4, pp. 237–250, 1982.",
    "S. K. Dhall and C. L. Liu, \"On a real-time scheduling problem,\" Operations Research, vol. 26, no. 1, "
    "pp. 127–140, 1978.",
    "G. C. Buttazzo, Hard Real-Time Computing Systems: Predictable Scheduling Algorithms and Applications, 3rd ed. "
    "Springer, 2011.",
    "R. I. Davis and A. Burns, \"A survey of hard real-time scheduling for multiprocessor systems,\" ACM Computing "
    "Surveys, vol. 43, no. 4, 2011.",
    "Real-Time Systems, Unit 2 lecture notes and lab manual (Mini Projects), University of Rwanda, 2026.",
]
for i, r in enumerate(refs, 1):
    p = para(f"[{i}]  {r}", size=10, space_after=4)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)

heading("Appendix A – Reproducing the results", 1)
body("Clone the repository, install matplotlib, and run the scripts from the repository root. Each script prints its "
     "tables and writes its charts to figures/ and its text output to results/.")
code("""git clone https://github.com/hirwagilbert/lab--2-real-time-system-
cd lab--2-real-time-system-
pip install matplotlib
python mp21_run.py                                # Mini Project 2.1
python schedcheck.py tasksets/*.csv --out figures # Mini Project 2.2
python partition.py                               # Mini Project 2.3
python tests/test_all.py                          # hand-computed checks""")
body("Using the tool from Python with an in-script list:")
code("""from rts_sim import Task
from schedcheck import analyze, format_report
rep = analyze([Task("t1", 1, 4), Task("t2", 2, 6), Task("t3", 2, 12, D=3)])
print(format_report(rep))""")

p = para(space_after=4)
p.paragraph_format.space_before = Pt(14)
bottom_rule(p, "B8C7D3", 4)
para("Source code and results", bold=True, size=11, color=BLUE, space_after=2)
para("All code, task sets, charts and this report are available in the group's GitHub repository:", space_after=2)
para(REPO, bold=True, size=11, color=BLUE, space_after=0)

out ="report/RTS_Unit2_Lab_Report.docx"
doc.save(out)
json.dump([t for _, t in TOC_ENTRIES], open("report/toc_entries.json", "w"))
missing = [t for _, t in TOC_ENTRIES if t not in [h for _, h in HEADINGS]]
if missing:
    print("WARNING: TOC entries without heading:", missing)
print("saved", out)
